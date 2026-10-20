# API/utils/document_parser.py

import os
import asyncio
import ipaddress
import io
import re
import base64
import socket
from urllib.parse import urljoin, urlsplit
from typing import Dict, Any, Optional, Tuple
import httpx
from pypdf import PdfReader
import docx

MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_DOWNLOAD_REDIRECTS = 5


async def _validate_download_url(url: str) -> str:
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("La URL de descarga debe usar HTTP o HTTPS")
    if parsed.username or parsed.password:
        raise ValueError("No se permiten credenciales embebidas en la URL")

    hostname = parsed.hostname.rstrip(".").lower()
    if hostname in {"localhost", "localhost.localdomain"} or hostname.endswith(".local"):
        raise ValueError("No se permiten destinos locales")

    allowed_hosts = {
        host.strip().lower().lstrip(".")
        for host in os.getenv("ATTACHMENT_ALLOWED_HOSTS", "").split(",")
        if host.strip()
    }
    if allowed_hosts and not any(
        hostname == allowed or hostname.endswith(f".{allowed}")
        for allowed in allowed_hosts
    ):
        raise ValueError("El host de descarga no está permitido")
    if os.getenv("APP_ENV", "").strip().lower() in {"prod", "production"} and not allowed_hosts:
        raise ValueError("ATTACHMENT_ALLOWED_HOSTS debe configurarse en producción")

    try:
        literal_ip = ipaddress.ip_address(hostname)
        addresses = [literal_ip]
    except ValueError:
        port = parsed.port or (443 if parsed.scheme == "https" else 80)
        records = await asyncio.to_thread(
            socket.getaddrinfo, hostname, port, 0, socket.SOCK_STREAM
        )
        addresses = [ipaddress.ip_address(record[4][0]) for record in records]
    if not addresses or any(not address.is_global for address in addresses):
        raise ValueError("No se permiten destinos de red privados o reservados")
    return hostname


async def fetch_url_bytes(url: str, timeout: float = 30.0) -> Tuple[bytes, str]:
    """
    Descarga el contenido binario desde una URL remota.
    Retorna los bytes y el nombre/tipo sugerido en los headers.
    """
    current_url = url
    async with httpx.AsyncClient(
        follow_redirects=False, timeout=timeout, trust_env=False
    ) as client:
        for redirect_count in range(MAX_DOWNLOAD_REDIRECTS + 1):
            await _validate_download_url(current_url)
            async with client.stream("GET", current_url) as response:
                if 300 <= response.status_code < 400:
                    location = response.headers.get("location")
                    if not location:
                        raise ValueError("La redirección de descarga no incluye destino")
                    if redirect_count == MAX_DOWNLOAD_REDIRECTS:
                        raise ValueError("Se excedió el máximo de redirecciones de descarga")
                    next_url = urljoin(current_url, location)
                    if (
                        urlsplit(current_url).scheme == "https"
                        and urlsplit(next_url).scheme != "https"
                    ):
                        raise ValueError("No se permiten redirecciones HTTPS a HTTP")
                    current_url = next_url
                    continue
                response.raise_for_status()
                content_length = response.headers.get("content-length")
                if content_length and int(content_length) > MAX_DOWNLOAD_BYTES:
                    raise ValueError("El archivo excede el máximo de 25 MB")
                chunks = []
                size = 0
                async for chunk in response.aiter_bytes():
                    size += len(chunk)
                    if size > MAX_DOWNLOAD_BYTES:
                        raise ValueError("El archivo excede el máximo de 25 MB")
                    chunks.append(chunk)

                content_type = response.headers.get("content-type", "").lower()
                content_disposition = response.headers.get("content-disposition", "")
                break
        else:
            raise ValueError("No se pudo completar la descarga")
        
        # Intentar extraer nombre de archivo si viene en Content-Disposition
        filename = ""
        if "filename=" in content_disposition:
            match = re.search(r'filename=["\']?([^"\';]+)["\']?', content_disposition)
            if match:
                filename = match.group(1)
                
        return b"".join(chunks), filename or content_type


def detect_file_type(data: bytes, hint: str = "") -> str:
    """
    Detecta el tipo de documento analizando magic bytes o extensiones/pistas.
    Retorna: 'pdf', 'docx', o 'text'.
    """
    hint_lower = hint.lower()
    
    # 1. Magic bytes PDF
    if data.startswith(b"%PDF"):
        return "pdf"
    
    # 2. DOCX es un archivo ZIP estándar (PK\x03\x04) que contiene estructuras de Word
    if data.startswith(b"PK\x03\x04"):
        if (
            b"word/" in data[:4000]
            or b"[Content_Types].xml" in data[:2000]
            or ".docx" in hint_lower
            or "word" in hint_lower
            or "officedocument" in hint_lower
            or not hint
        ):
            return "docx"
            
    # Pistas por extensión o MIME type
    if hint_lower.endswith(".pdf") or "application/pdf" in hint_lower:
        return "pdf"
    if hint_lower.endswith(".docx") or "wordprocessingml" in hint_lower or hint_lower.endswith(".doc"):
        return "docx"
    
    return "text"


def extract_text_from_pdf(data: bytes) -> Dict[str, Any]:
    """Extrae texto estructurado y metadatos de un archivo PDF usando pypdf."""
    stream = io.BytesIO(data)
    reader = PdfReader(stream)
    
    pages_text = []
    for i, page in enumerate(reader.pages):
        page_content = page.extract_text() or ""
        pages_text.append(page_content)
        
    full_text = "\n\n".join(pages_text).strip()
    words = len(full_text.split())
    
    return {
        "text": full_text,
        "metadata": {
            "page_count": len(reader.pages),
            "word_count": words,
            "char_count": len(full_text)
        }
    }


def extract_text_from_docx(data: bytes) -> Dict[str, Any]:
    """Extrae texto, párrafos y tablas de un archivo Word (.docx) usando python-docx."""
    stream = io.BytesIO(data)
    doc = docx.Document(stream)
    
    content_parts = []
    
    # Extraer párrafos
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            content_parts.append(text)
            
    # Extraer contenido de tablas (muy relevante para rúbricas)
    table_count = len(doc.tables)
    for table_idx, table in enumerate(doc.tables, start=1):
        content_parts.append(f"\n--- [Tabla de Rúbrica #{table_idx}] ---")
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            content_parts.append(" | ".join(row_cells))
            
    full_text = "\n".join(content_parts).strip()
    words = len(full_text.split())
    
    return {
        "text": full_text,
        "metadata": {
            "paragraph_count": len(doc.paragraphs),
            "table_count": table_count,
            "word_count": words,
            "char_count": len(full_text)
        }
    }


def is_base64_string(s: str) -> bool:
    """Comprueba si una cadena es una representación Base64 válida."""
    if len(s) < 32 or " " in s:
        return False
    # Remover encabezados data URI si los tiene
    if s.startswith("data:") and ";base64," in s:
        return True
    try:
        # Longitud múltiplo de 4 y caracteres base64
        if len(s) % 4 != 0:
            return False
        decoded = base64.b64decode(s, validate=True)
        return len(decoded) > 0
    except Exception:
        return False


async def parse_document_input(source: str, hint: str = "") -> Dict[str, Any]:
    """
    Función universal para procesar cualquier entrada de documento:
    - URL (http/https)
    - Ruta de archivo local en disco
    - Cadena Base64
    - Texto plano directo (sin conversión)
    """
    if not source or not source.strip():
        return {
            "status": "error",
            "error": "La entrada del documento está vacía.",
            "text": "",
            "metadata": {}
        }
        
    source_clean = source.strip()
    source_type = "raw_text"
    raw_bytes: Optional[bytes] = None
    file_hint = hint
    
    try:
        # 1. Caso: URL de descarga
        if source_clean.startswith("http://") or source_clean.startswith("https://"):
            source_type = "url"
            raw_bytes, url_hint = await fetch_url_bytes(source_clean)
            file_hint = url_hint or source_clean.split("?")[0].split("/")[-1] or hint

        # 2. Caso: Archivo local en disco
        elif os.path.exists(source_clean) and os.path.isfile(source_clean):
            if os.getenv("APP_ENV", "").strip().lower() in {"prod", "production"}:
                raise ValueError("Las rutas locales no están habilitadas en producción")
            source_type = "file"
            with open(source_clean, "rb") as f:
                raw_bytes = f.read()
            file_hint = os.path.basename(source_clean)

        # 3. Caso: Cadena codificada en Base64
        elif is_base64_string(source_clean):
            source_type = "base64"
            base64_payload = source_clean
            if "base64," in source_clean:
                prefix = source_clean.split("base64,")[0]
                file_hint = prefix or hint
                base64_payload = source_clean.split("base64,")[1]
            raw_bytes = base64.b64decode(base64_payload)

        # 4. Caso: Texto plano directo (no binario)
        else:
            source_type = "raw_text"
            words = len(source_clean.split())
            return {
                "status": "success",
                "source_type": source_type,
                "file_type": "text",
                "text": source_clean,
                "metadata": {
                    "word_count": words,
                    "char_count": len(source_clean),
                    "source": "plain_text"
                }
            }

        # Si tenemos bytes, determinamos el formato y extraemos
        file_type = detect_file_type(raw_bytes, file_hint)
        
        if file_type == "pdf":
            extracted = extract_text_from_pdf(raw_bytes)
        elif file_type == "docx":
            extracted = extract_text_from_docx(raw_bytes)
        else:
            # Texto plano decodificado de bytes
            try:
                text_content = raw_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text_content = raw_bytes.decode("latin-1", errors="replace")
            extracted = {
                "text": text_content.strip(),
                "metadata": {
                    "word_count": len(text_content.split()),
                    "char_count": len(text_content)
                }
            }
            
        extracted["metadata"]["source_type"] = source_type
        extracted["metadata"]["detected_type"] = file_type
        extracted["metadata"]["file_hint"] = file_hint
        
        return {
            "status": "success",
            "source_type": source_type,
            "file_type": file_type,
            "text": extracted["text"],
            "metadata": extracted["metadata"]
        }

    except Exception as e:
        return {
            "status": "error",
            "source_type": source_type,
            "error": f"Error al procesar el documento: {str(e)}",
            "text": "",
            "metadata": {}
        }
